from datetime import timedelta
import io
import pytest
from backend.analytics import sensor_health,is_expected,label_window,metrics
from backend.ingestion import number,parse_date,read_rows,ingest_events,contains_pii
from backend.context import ContextBuilder
from backend.db import Record,now
from backend.prediction_providers.rules import RulesProvider
from backend.prediction_providers.http import ProviderRouter
from backend.services import run_predictions
from tools.emulator.scenarios import inject_scenario

def test_russian_parser():
    assert number('3 008 235 019')==3008235019
    assert number('25,00')==25
    assert parse_date('19.10.2026 12:15').hour==9
    content='ИД записи журнала;ИД канала данных;Текущее значение;Дата записи\n3 008 235 019;56 682;25,00;19.10.2026 12:15'.encode('cp1251')
    assert read_rows(content,'sample.csv')[0]['value']=='25,00'
    with pytest.raises(ValueError):
        number('nan')

def test_xlsx_merged_multiple():
    from openpyxl import Workbook
    wb=Workbook()
    for sheet in [wb.active,wb.create_sheet('Второй')]:
        sheet.append(['Отчёт']);sheet.merge_cells('A1:D1')
        sheet.append(['id','channel_id','value','ts']);sheet.append(['a',12,'25,4','19.10.2026 12:15'])
    stream=io.BytesIO();wb.save(stream)
    assert len(read_rows(stream.getvalue(),'sample.xlsx'))==2

def test_pii_idempotency(populated):
    rows=[{'id':'external-1','channel_id':12,'value':'25,00','ts':'19.10.2026 12:15'}]
    assert ingest_events(populated,rows)['accepted']==1
    assert ingest_events(populated,rows)['duplicates']==1
    assert contains_pii({'comment':'mail@example.com'})
    with pytest.raises(ValueError):
        ingest_events(populated,[{**rows[0],'comment':'+7 999 123-45-67'}])
    assert ingest_events(populated,[{'channel_id':999,'ts':'bad','value':2}])['errors']


def test_ml_journal_rows_keep_source_alarm(populated):
    from sqlalchemy import select
    from backend.db import Channel, Event

    content = ('"ид_события","ид_канала_данных","дата","время","тревожное","значение_датчика"\n'
               '4524243389,120473,"2026-08-01","03:09:27",false,"28"\n'
               '4524253385,120475,"2026-08-01","03:19:55",false,"29"\n').encode()
    rows = read_rows(content, 'journal.csv')
    assert ingest_events(populated, rows)['errors'][0]['message'] == 'Неизвестный канал'

    populated.add(Channel(id=120473, object_id='obj-000001', type_id=5, type_code='movement'))
    populated.add(Channel(id=120475, object_id='obj-000001', type_id=5, type_code='movement'))
    populated.flush()
    result = ingest_events(populated, rows)
    assert result['accepted'] == 2 and not result['errors']
    events = populated.scalars(select(Event).where(Event.id.in_(['4524243389', '4524253385'])).order_by(Event.id)).all()
    assert [event.ts.hour for event in events] == [0, 0]
    assert [event.value for event in events] == [28, 29]
    assert all(event.severity == 'normal' for event in events)
    assert all(event.raw['тревожное'] == 'false' for event in events)
    text_state = {**rows[0], 'id': 'text-state', 'value': 'Неопределен', 'тревожное': 'true'}
    result = ingest_events(populated, [text_state])
    assert result['accepted'] == 1 and not result['errors']
    state = populated.get(Event, 'text-state')
    assert state.value is None
    assert state.raw['value'] == 'Неопределен'
    assert state.severity == 'warning'
    context = ContextBuilder(populated).build(['obj-000001'], parse_date('2026-08-01 03:10:00'))
    movement = next(c for c in context.objects[0].channels if c.channel_id == 120473)
    assert any(value is None for _, value in movement.series.points)

def test_health():
    t=now();points=[(t-timedelta(minutes=5*i),20) for i in range(12)]
    assert sensor_health(points,t)['flatline']
    assert not sensor_health(points,t,'smoke')['flatline']
    assert sensor_health([],t)['dropout_rate']==1
    assert sensor_health([(t-timedelta(minutes=5*i),i%2) for i in range(12)],t,'movement')['chatter']
    assert sensor_health(points,t,neighbor_mean=0)['drift']>8

def test_works():
    t=now();w=[{'from':t.isoformat(),'to':(t+timedelta(hours=1)).isoformat(),'status':'approved'}]
    assert is_expected(t-timedelta(minutes=10),w)
    assert not is_expected(t-timedelta(minutes=16),w)
    assert not is_expected(t,[{**w[0],'status':'cancelled'}])

def test_metrics_labels():
    t=now()
    assert label_window(t,24,[t+timedelta(hours=10)],t+timedelta(hours=25))==1
    assert label_window(t,24,[t],t+timedelta(hours=25))==0
    assert label_window(t,24,[],t+timedelta(hours=20)) is None
    m=metrics([1,0,1,0],[.9,.2,.8,.1])
    assert m['precision']==m['recall']==m['f1']==m['pr_auc']==1
    assert metrics([1,0],[.5,.5])['pr_auc']==.5
    assert metrics([],[])['brier'] is None

def test_no_future_leakage(populated):
    t=now()
    populated.add(Record(id='future',kind='weather',object_id='obj-000001',ts=t+timedelta(hours=2),available_at=t+timedelta(hours=1),data={'kind':'forecast','ts':(t+timedelta(hours=2)).isoformat(),'published_at':(t+timedelta(hours=1)).isoformat(),'precip_mm':100}))
    populated.flush()
    assert not ContextBuilder(populated).build(['obj-000001'],t).objects[0].weather['forecast']

def test_rules_scenarios_notifications(populated):
    for kind,obj in [('fire','obj-000001'),('flood','obj-000002')]:
        inject_scenario(populated,kind)
        result=run_predictions(populated,RulesProvider(),[obj])
        assert next(p for p in result if p['incident_type']==kind)['probability']>.75
    from sqlalchemy import select
    notifications=populated.scalars(select(Record).where(Record.kind=='notification')).all()
    assert notifications
    run_predictions(populated,RulesProvider())
    assert len(populated.scalars(select(Record).where(Record.kind=='notification')).all())==len(notifications)
    request=ContextBuilder(populated).build(['obj-000002'],now())
    baseline=RulesProvider().predict(request)
    request.options.what_if={'weather.forecast.precip_mm_delta':10}
    assert RulesProvider().predict(request).predictions[1].probability>=baseline.predictions[1].probability

def test_fallback_recovery(populated,monkeypatch):
    import httpx
    inject_scenario(populated,'fire');request=ContextBuilder(populated).build(['obj-000001'],now())
    router=ProviderRouter();router.mode='http'
    def fail(r):
        raise httpx.ConnectError('unavailable')
    monkeypatch.setattr(router.remote,'predict',fail)
    response=router.predict(request)
    assert response.model.kind=='rules' and router.health()['degraded']
    assert router.predict(request)==response
    router.open_until=0
    monkeypatch.setattr(router.remote,'predict',RulesProvider().predict)
    request.request_id='recovery';router.predict(request)
    assert not router.degraded

def test_readonly_adapters(monkeypatch):
    import httpx
    from backend.adapters import SmvuAdapter,EquipmentRegistryAdapter,OdsJournalAdapter,WorkOrdersAdapter
    called=[]
    def get(self,url,**kwargs):
        called.append('GET');return httpx.Response(200,json={'items':[{'id':'a'}]},request=httpx.Request('GET',url))
    monkeypatch.setattr(httpx.Client,'get',get)
    for cls in [SmvuAdapter,EquipmentRegistryAdapter,OdsJournalAdapter,WorkOrdersAdapter]:
        assert cls(url='https://example.test/data').fetch()==[{'id':'a'}]
        assert not hasattr(cls,'write')
    assert called==['GET']*4
