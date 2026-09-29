"""Импорт источников с отчётом, нормализацией и идемпотентностью."""
import csv
import hashlib
import io
import math
import re
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from sqlalchemy import select
from backend.analytics import is_expected
from backend.db import Channel, Event, Object, Record, utc
from backend.settings import config

TYPES = {2: 'contact-unlock-norm', 3: 'switch', 4: 'smoke', 5: 'movement', 6: 'gas', 7: 'pump', 8: 'fan', 9: 'phase', 12: 'temperature'}
CUSTOMER_TYPE_CODES = {
    'КД Дверь': 'contact-unlock-norm', 'КД Люк': 'contact-unlock-norm',
    'КД АВ': 'contact-unlock-norm', 'Переключатель': 'switch',
    'Датчик дыма': 'smoke', 'Датчик движения': 'movement',
    'Газовый датчик': 'gas', 'Состояние насоса': 'pump',
    'Состояние вентилятора': 'fan', 'Состояние фазы': 'phase',
    'Датчик температуры': 'temperature',
}
ALIASES = {'ИД записи журнала': 'id', 'ИД канала данных': 'channel_id', 'ИД типа канала': 'type_id', 'ИД типа канала данных': 'type_id', 'Текущее значение': 'value', 'Дата записи': 'ts', 'Тег в дереве объектов': 'tag', 'ид_события': 'id', 'ид_канала_данных': 'channel_id', 'значение_датчика': 'value'}
PII = [r'[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}', r'(?<!\d)(?:\+7|8)[\s(-]*\d{3}[\s)-]*\d{3}[\s-]*\d{2}[\s-]*\d{2}(?!\d)', r'\b\d{4}\s+\d{6}\b', r'\b[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:вич|вна)\b']


def contains_pii(value) -> bool:
    if isinstance(value, dict):
        return any(contains_pii(v) for k, v in value.items() if k not in ('id', 'channel_id', 'type_id', 'object_id', 'ts', 'value'))
    if isinstance(value, list):
        return any(contains_pii(v) for v in value)
    return isinstance(value, str) and any(re.search(p, value) for p in PII)


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        except ValueError:
            parsed = datetime.strptime(str(value), '%d.%m.%Y %H:%M')
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo('Europe/Moscow'))
    return utc(parsed)


def number(value) -> float:
    result = float(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
    if not math.isfinite(result):
        raise ValueError('Число не конечное')
    return result


def source_alarm(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in ('true', 'false'):
        return value.strip().lower() == 'true'
    raise ValueError('Неверное значение тревожное')


def read_rows(content: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith('.xlsx'):
        wb = load_workbook(io.BytesIO(content), data_only=True)
        result = []
        for sheet in wb:
            for merged in list(sheet.merged_cells.ranges):
                value = sheet.cell(merged.min_row, merged.min_col).value
                sheet.unmerge_cells(str(merged))
                for row in sheet.iter_rows(min_row=merged.min_row, max_row=merged.max_row, min_col=merged.min_col, max_col=merged.max_col):
                    for cell in row:
                        cell.value = value
            headers = None
            for row in sheet.iter_rows(values_only=True):
                if not any(v is not None for v in row):
                    continue
                normalized = [ALIASES.get(str(v).strip(), str(v).strip()) for v in row]
                if headers is None:
                    if any(k in normalized for k in ('channel_id', 'object_id', 'id', 'tag')):
                        headers = normalized
                    continue
                result.append(dict(zip(headers, row)))
        if not result:
            raise ValueError('Не найдены строки и заголовки XLSX')
        return result
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('cp1251')
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;')
    except csv.Error:
        dialect = csv.excel
    return [{ALIASES.get(str(k).strip(), str(k).strip()): v for k, v in row.items()} for row in csv.DictReader(io.StringIO(text), dialect=dialect)]


def ingest_events(session, rows: list[dict]) -> dict:
    result = dict(accepted=0, duplicates=0, errors=[], warnings=[])
    # Отклоняем весь пакет до записи, чтобы ПДн не попали даже в частичный импорт.
    if contains_pii(rows):
        raise ValueError('Обнаружены возможные персональные данные; пакет отклонён')
    works = session.scalars(select(Record).where(Record.kind == 'work')).all()
    work_map = {}
    for w in works:
        work_map.setdefault(w.object_id, []).append(w.data)
    for index, row in enumerate(rows, 1):
        try:
            source = row
            row = {ALIASES.get(key, key): value for key, value in source.items()}
            if 'ts' not in row and 'дата' in row and 'время' in row:
                row['ts'] = f"{row['дата']} {row['время']}"
            channel_id = int(number(row['channel_id']))
            channel = session.get(Channel, channel_id)
            if channel is None:
                raise ValueError('Неизвестный канал')
            if row.get('type_id') and int(number(row['type_id'])) != channel.type_id:
                raise ValueError('Тип канала не совпадает с реестром')
            ts = parse_date(row['ts'])
            raw_value = row['value']
            try:
                value = number(raw_value)
            except (ValueError, TypeError):
                value = None
            if value is None and 'тревожное' not in row:
                raise ValueError('Для текстового значения требуется тревожное')
            identity = str(row.get('id') or hashlib.sha256(f'{channel_id}:{ts.isoformat()}:{raw_value}'.encode()).hexdigest()).replace(' ', '')
            if session.get(Event, identity):
                result['duplicates'] += 1
                continue
            expected = is_expected(ts, work_map.get(channel.object_id, []), config('settings')['planned_work_tolerance_minutes'])
            if 'тревожное' in row:
                alarm = source_alarm(row['тревожное'])
            elif channel.type_code == 'temperature':
                alarm = value > 45
            elif channel.type_code in ('phase', 'fan'):
                alarm = value < 0.5
            else:
                alarm = value > 0.5
            if value is None:
                event = str(raw_value)
            elif alarm and channel.type_code == 'movement':
                event = 'Обнаружено движение'
            else:
                event = 'Предупреждение' if alarm else 'Норма'
            session.add(Event(id=identity, channel_id=channel_id, object_id=channel.object_id, ts=ts, value=value, expected=expected, event=event, severity='warning' if alarm else 'normal', raw={k: str(v) for k, v in source.items()}))
            if channel.type_code == 'temperature' and value is not None and not -50 <= value <= 150:
                result['warnings'].append({'row': index, 'message': 'Выброс температуры сохранён для проверки'})
            session.flush()
            result['accepted'] += 1
        except (ValueError, KeyError, TypeError) as exc:
            result['errors'].append({'row': index, 'message': str(exc)})
    return result


def ingest_registry(session, rows: list[dict], kind: str) -> dict:
    if contains_pii(rows):
        raise ValueError('Обнаружены персональные данные')
    count, errors = 0, []
    for index, row in enumerate(rows, 1):
        try:
            if kind == 'objects' and 'ид_объект' in row:
                identity = str(row['ид_объект']).strip()
                existing = session.get(Object, identity)
                meta = dict(existing.meta or {}) if existing else {}
                meta.update({
                    'source': 'customer_registry',
                    'object_kind': str(row['вид_объекта']).strip(),
                    'parent_id': str(row['родитель']).strip(),
                    'hierarchy_level': str(row['иерархия_уровень']).strip(),
                    'location_status': 'provided' if existing and existing.lat is not None and existing.lon is not None else 'unavailable',
                })
                session.merge(Object(
                    id=identity, name=str(row['диспетчерское_название_объекта']).strip(),
                    tag=existing.tag if existing else f'customer-object:{identity}',
                    lat=existing.lat if existing else None, lon=existing.lon if existing else None,
                    meta=meta,
                ))
            elif kind == 'channels' and ('ид_канала_данных' in row or 'channel_id' in row) and 'тип_датчика' in row:
                identity = int(number(row.get('channel_id', row.get('ид_канала_данных'))))
                object_id = str(row['ид_объект']).strip()
                obj = session.get(Object, object_id)
                if obj is None:
                    raise ValueError('Неизвестный объект; сначала загрузите справочник объектов')
                existing = session.get(Channel, identity)
                if existing and existing.object_id != object_id:
                    raise ValueError('Канал уже привязан к другому объекту')
                passport_id = f'ml-passport:{identity}'
                passport = session.get(Record, passport_id)
                if passport and passport.data.get('source') != 'customer_registry':
                    raise ValueError('ML-паспорт канала задан вручную')
                sensor = str(row['тип_датчика']).strip()
                code = CUSTOMER_TYPE_CODES.get(sensor, 'unclassified')
                type_id = next((key for key, value in TYPES.items() if value == code), 0)
                session.merge(Channel(id=identity, object_id=object_id, type_id=type_id, type_code=code))
                session.merge(Record(
                    id=passport_id, kind='ml_channel_passport', object_id=object_id,
                    data={
                        'sensor_type': sensor,
                        'engineering_system': str(row['тип_инж_системы']).strip(),
                        'object_kind': obj.meta['object_kind'],
                        'parent': obj.meta['parent_id'],
                        'synthetic': False, 'source': 'customer_registry',
                        'sensor_name': str(row['название_датчика']).strip(),
                        'engineering_tag': str(row['тег_инженерной_системы']).strip(),
                    },
                ))
            elif kind == 'objects':
                session.merge(Object(id=str(row['id']), name=str(row['name']), tag=str(row['tag']), lat=number(row['lat']), lon=number(row['lon']), meta=row.get('meta', {}) if isinstance(row.get('meta', {}), dict) else {}))
            elif kind == 'channels':
                type_id = int(number(row['type_id']))
                if type_id not in TYPES or not session.get(Object, str(row['object_id'])):
                    raise ValueError('Неизвестный тип или объект')
                session.merge(Channel(id=int(number(row.get('channel_id', row.get('id')))), object_id=str(row['object_id']), type_id=type_id, type_code=TYPES[type_id]))
            else:
                raise ValueError('Неизвестный вид справочника')
            session.flush()
            count += 1
        except (ValueError, KeyError, TypeError) as exc:
            errors.append({'row': index, 'message': str(exc)})
    return {'accepted': count, 'errors': errors}
